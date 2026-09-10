# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import copy

NAVY = RGBColor(0x1E, 0x27, 0x61)
INK = RGBColor(0x2A, 0x33, 0x44)
INK_SOFT = RGBColor(0x4A, 0x5B, 0x72)
ACCENT = RGBColor(0x2F, 0x6F, 0xB3)
LINE = RGBColor(0xDB, 0xE3, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFTBG = RGBColor(0xF2, 0xF5, 0xFC)
GOOD = RGBColor(0x1E, 0x7B, 0x4D)
BAD = RGBColor(0xB0, 0x3A, 0x2E)

CAP = r"C:\dev\workspace-igm\advanced-cloud-developer\실습산출물\3차시\캡처"
MOCKCAP = CAP
DIAG = r"C:\dev\workspace-igm\advanced-cloud-developer\scratch_pptx\diagrams"

FONT_BODY = "Calibri"

def _clear_extra_runs(paragraph):
    for r in list(paragraph.runs)[1:]:
        r._r.getparent().remove(r._r)


def _clear_extra_paragraphs(text_frame):
    for p in list(text_frame.paragraphs)[1:]:
        p._p.getparent().remove(p._p)


def set_section_title(slide, text):
    ph = slide.placeholders[0]
    tf = ph.text_frame
    _clear_extra_paragraphs(tf)
    _clear_extra_runs(tf.paragraphs[0])
    tf.paragraphs[0].runs[0].text = text


def find_title_box(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() in ("제목", ""):
            if sh.name.startswith("Google Shape") or sh.name.startswith("TextBox"):
                return sh
    return None


def add_title(slide, text, size=30):
    box = find_title_box(slide)
    if box is not None:
        tf = box.text_frame
        _clear_extra_paragraphs(tf)
        _clear_extra_runs(tf.paragraphs[0])
        if not tf.paragraphs[0].runs:
            tf.paragraphs[0].add_run()
        run = tf.paragraphs[0].runs[0]
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = True
        run.font.name = FONT_BODY
        run.font.color.rgb = NAVY
        return box
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.28), Inches(12.33), Inches(0.9))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = FONT_BODY
    r.font.color.rgb = NAVY
    return box


def add_kicker(slide, text):
    box = slide.shapes.add_textbox(Inches(0.52), Inches(0.02), Inches(8), Inches(0.3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.name = FONT_BODY
    r.font.color.rgb = ACCENT
    return box


def add_rule_space(slide):
    pass


def add_bullets(slide, items, left, top, width, height, size=15, color=INK, bold_first=False,
                 line_spacing=1.15, space_after=8, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        if isinstance(item, tuple):
            head, rest = item
            r1 = p.add_run()
            r1.text = head
            r1.font.bold = True
            r1.font.size = Pt(size)
            r1.font.name = FONT_BODY
            r1.font.color.rgb = NAVY
            r2 = p.add_run()
            r2.text = rest
            r2.font.size = Pt(size)
            r2.font.name = FONT_BODY
            r2.font.color.rgb = color
        else:
            r = p.add_run()
            r.text = "•  " + item
            r.font.size = Pt(size)
            r.font.name = FONT_BODY
            r.font.color.rgb = color
    return box


def add_card(slide, left, top, width, height, fill=SOFTBG, line=LINE, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    if radius:
        try:
            shp.adjustments[0] = 0.06
        except Exception:
            pass
    return shp


def add_stat(slide, left, top, width, height, number, label, num_color=NAVY):
    card = add_card(slide, left, top, width, height, fill=WHITE, line=LINE)
    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = number
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.name = FONT_BODY
    r.font.color.rgb = num_color
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = label
    r2.font.size = Pt(11)
    r2.font.name = FONT_BODY
    r2.font.color.rgb = INK_SOFT
    return card


def style_table(table, header_fill=NAVY, header_font=WHITE, font_size=11, col_widths=None,
                 body_fill=WHITE, alt_fill=SOFTBG):
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = Inches(w)
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if r_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_fill
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = body_fill if r_idx % 2 == 1 else alt_fill
            for p in cell.text_frame.paragraphs:
                p.line_spacing = 1.05
                for run in p.runs:
                    run.font.size = Pt(font_size)
                    run.font.name = FONT_BODY
                    if r_idx == 0:
                        run.font.bold = True
                        run.font.color.rgb = header_font
                    else:
                        run.font.color.rgb = INK


def add_table(slide, left, top, width, height, data, col_widths=None, font_size=11):
    rows = len(data)
    cols = len(data[0])
    gframe = slide.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(width), Inches(height))
    table = gframe.table
    for r_idx, row in enumerate(data):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text_frame.clear()
            p = cell.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = str(val)
    style_table(table, col_widths=col_widths, font_size=font_size)
    return gframe


def add_pic_fit(slide, path, left, top, max_width, max_height, caption=None, align="center"):
    """Place an image so it fits within (max_width, max_height) box, preserving aspect ratio."""
    from PIL import Image
    with Image.open(path) as im:
        px_w, px_h = im.size
    ratio = px_w / px_h
    box_ratio = max_width / max_height
    if ratio >= box_ratio:
        w = max_width
        h = w / ratio
    else:
        h = max_height
        w = h * ratio
    if align == "center":
        left = left + (max_width - w) / 2
    pic = slide.shapes.add_picture(path, Inches(left), Inches(top), width=Inches(w), height=Inches(h))
    if caption:
        cap_box = slide.shapes.add_textbox(Inches(left), Inches(top + h + 0.06), Inches(w), Inches(0.32))
        tf = cap_box.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = caption
        r.font.size = Pt(10.5)
        r.font.italic = True
        r.font.name = FONT_BODY
        r.font.color.rgb = INK_SOFT
    return pic


def add_pic_framed(slide, path, left, top, width=None, height=None, caption=None):
    pic = slide.shapes.add_picture(path, Inches(left), Inches(top), width=Inches(width) if width else None,
                                    height=Inches(height) if height else None)
    if caption:
        cap_box = slide.shapes.add_textbox(Inches(left), Inches(top + (height or pic.height / 914400) + 0.05),
                                            Inches(width or pic.width / 914400), Inches(0.3))
        tf = cap_box.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = caption
        r.font.size = Pt(10.5)
        r.font.italic = True
        r.font.name = FONT_BODY
        r.font.color.rgb = INK_SOFT
    return pic


def build_pass1(prs):
    slides = prs.slides
    s1 = slides[0]
    for sh in s1.shapes:
        if sh.name == "TextBox 4":
            tf = sh.text_frame
            _clear_extra_paragraphs(tf)
            _clear_extra_runs(tf.paragraphs[0])
            tf.paragraphs[0].runs[0].text = "레포 인사이트"
            p2 = tf.add_paragraph()
            r2 = p2.add_run()
            r2.text = "GitHub 레포 분석 → Notion 포트폴리오 자동화 도구"
            r2.font.size = Pt(20)
            r2.font.bold = False
            r2.font.color.rgb = WHITE
            r2.font.name = FONT_BODY
        elif sh.name == "TextBox 8":
            tf = sh.text_frame
            _clear_extra_paragraphs(tf)
            _clear_extra_runs(tf.paragraphs[0])
            tf.paragraphs[0].runs[0].text = "Advanced Cloud Developer 심화 과정 · 3차시 실습 결과보고"

    SECTION_MAP = {
        1: "1. 프로젝트 개요",
        3: "2. 개발 환경 및 요구사항 분석",
        7: "3. 앱 설계",
        14: "4. 테스트 결과",
        16: "5. 배포",
        19: "6. 과정 회고",
        21: "7. 시연",
    }
    for idx, text in SECTION_MAP.items():
        set_section_title(slides[idx], text)
    print("pass1 (title + sections) applied")


if __name__ == "__main__":
    prs = Presentation("report.pptx")
    build_pass1(prs)
    prs.save("report.pptx")
    print("saved")
